# Copyright (c) 2023-present Plane Software, Inc. and contributors
# SPDX-License-Identifier: AGPL-3.0-only
# See the LICENSE file for details.

"""CSV and XLSX rendering for report exports."""

import csv
import io
import re
from typing import Any, Dict, List, Sequence, Tuple
from urllib.parse import quote

from django.http import HttpResponse
from django.utils import timezone
from django.utils.text import slugify
from openpyxl import Workbook

from plane.utils.csv_utils import sanitize_csv_row

CSV_CONTENT_TYPE = "text/csv; charset=utf-8"
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EXPORT_FORMATS = ("csv", "xlsx")

_SHEET_INVALID = re.compile(r"[\[\]:*?/\\]")

# A sheet is (title, headers, rows) where rows are dicts keyed by header.
Sheet = Tuple[str, Sequence[str], List[Dict[str, Any]]]


def _cell(value: Any) -> Any:
    if isinstance(value, (list, tuple, set)):
        return ", ".join(str(v) for v in value if v not in (None, ""))
    if value is None:
        return ""
    return value


def render_csv(headers: Sequence[str], rows: List[Dict[str, Any]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(sanitize_csv_row(list(headers)))
    for row in rows:
        writer.writerow(sanitize_csv_row([_cell(row.get(h)) for h in headers]))
    # BOM so spreadsheet apps detect UTF-8 (non-Latin names, accents)
    return ("﻿" + buffer.getvalue()).encode("utf-8")


def safe_sheet_title(title: str, used: set) -> str:
    base = _SHEET_INVALID.sub(" ", title or "Sheet").strip()[:31] or "Sheet"
    candidate, n = base, 2
    while candidate.lower() in used:
        suffix = f" ({n})"
        candidate = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


def render_xlsx(sheets: List[Sheet]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    used: set = set()
    for title, headers, rows in sheets:
        sheet = workbook.create_sheet(safe_sheet_title(title, used))
        sheet.append(sanitize_csv_row(list(headers)))
        for row in rows:
            sheet.append(sanitize_csv_row([_cell(row.get(h)) for h in headers]))
    if not workbook.worksheets:
        workbook.create_sheet("Sheet")
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_filename(title: str, extension: str) -> str:
    stem = slugify(title or "", allow_unicode=False) or "report"
    return f"{stem[:80]}-{timezone.localdate().strftime('%Y%m%d')}.{extension}"


def file_response(content: bytes, title: str, export_format: str) -> HttpResponse:
    content_type = XLSX_CONTENT_TYPE if export_format == "xlsx" else CSV_CONTENT_TYPE
    filename = export_filename(title, export_format)
    response = HttpResponse(content, content_type=content_type)
    response["Content-Disposition"] = f"attachment; filename=\"{filename}\"; filename*=UTF-8''{quote(filename)}"
    response["Access-Control-Expose-Headers"] = "Content-Disposition"
    return response


def render_export(export_format: str, title: str, sheets: List[Sheet]) -> HttpResponse:
    """CSV uses the first sheet only; XLSX includes every sheet."""
    if export_format == "xlsx":
        return file_response(render_xlsx(sheets), title, "xlsx")
    _, headers, rows = sheets[0] if sheets else ("", [], [])
    return file_response(render_csv(headers, rows), title, "csv")
