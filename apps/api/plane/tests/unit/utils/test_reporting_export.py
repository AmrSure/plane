# Copyright (c) 2023-present Plane Software, Inc. and contributors
# SPDX-License-Identifier: AGPL-3.0-only
# See the LICENSE file for details.

"""Unit tests for report CSV/XLSX rendering."""

import io

import pytest
from openpyxl import load_workbook

from plane.utils.reporting import to_export_rows
from plane.utils.reporting.exporting import render_csv, render_export, render_xlsx, safe_sheet_title


@pytest.mark.unit
class TestReportExport:
    def test_csv_uses_given_headers_and_sanitizes_formulas(self):
        content = render_csv(["Name", "Total"], [{"Name": "=HYPERLINK(1)", "Total": 3}]).decode("utf-8")
        assert content.startswith("﻿Name,Total")
        assert "'=HYPERLINK(1)" in content

    def test_xlsx_has_sheets_and_sanitized_cells(self):
        content = render_xlsx(
            [
                ("Summary", ["Metric", "Value"], [{"Metric": "Scope", "Value": 4}]),
                ("Items: open/closed", ["Name"], [{"Name": "+cmd"}]),
            ]
        )
        workbook = load_workbook(io.BytesIO(content))
        assert workbook.sheetnames == ["Summary", "Items  open closed"]
        assert workbook["Items  open closed"]["A2"].value == "'+cmd"

    def test_sheet_titles_are_unique_and_short(self):
        used = set()
        first = safe_sheet_title("A" * 40, used)
        second = safe_sheet_title("A" * 40, used)
        assert len(first) == 31 and len(second) <= 31 and first != second

    def test_response_headers(self):
        response = render_export("csv", "Sprint report", [("S", ["A"], [{"A": 1}])])
        assert response["Content-Type"].startswith("text/csv")
        assert 'filename="sprint-report-' in response["Content-Disposition"]
        assert "filename*=UTF-8''" in response["Content-Disposition"]

    def test_pivot_rows(self):
        chart = {
            "data": [{"key": "high", "name": "high", "count": 3, "u1": 2, "none": 1}],
            "schema": {"u1": "Sara", "none": "none"},
        }
        headers, rows = to_export_rows(chart, "Priority", "Total", {"high": "High", "none": "Unassigned"})
        assert headers == ["Priority", "Sara", "Unassigned", "Total"]
        assert rows == [{"Priority": "High", "Sara": 2, "Unassigned": 1, "Total": 3}]
